"""
export_engine.py
================
Task 11: Export all results to CSV, Excel, JSON, HTML, and PDF.
Uses fpdf2 (pure Python, no system deps) for PDF generation.
"""

import json
import csv
import logging
from pathlib import Path
from typing import List, Dict, Any
from datetime import datetime

logger = logging.getLogger(__name__)


# ── Helpers ───────────────────────────────────────────────────────────────

def _flat(metrics_list):
    """Strip non-scalar fields for CSV/Excel export."""
    skip = {"confusion_matrix", "classification_report", "roc_curve", "pr_curve", "annotated_papers", "note"}
    return [{k: v for k, v in m.items() if k not in skip} for m in metrics_list]


# ── CSV ───────────────────────────────────────────────────────────────────

def export_csv(data: List[Dict], path: Path):
    if not data:
        return
    path.parent.mkdir(parents=True, exist_ok=True)
    with open(path, "w", newline="", encoding="utf-8") as f:
        writer = csv.DictWriter(f, fieldnames=data[0].keys())
        writer.writeheader()
        writer.writerows(data)
    logger.info(f"CSV -> {path}")


# ── Excel ─────────────────────────────────────────────────────────────────

def export_excel(sheets: Dict[str, List[Dict]], path: Path):
    try:
        import openpyxl
        path.parent.mkdir(parents=True, exist_ok=True)
        wb = openpyxl.Workbook()
        first = True
        for sheet_name, rows in sheets.items():
            if not rows:
                continue
            ws = wb.active if first else wb.create_sheet()
            ws.title = sheet_name[:31]
            first = False
            headers = list(rows[0].keys())
            ws.append(headers)
            for row in rows:
                ws.append([str(row.get(h, "")) for h in headers])
        wb.save(path)
        logger.info(f"Excel -> {path}")
    except ImportError:
        logger.warning("openpyxl not installed. Skipping Excel export. Run: pip install openpyxl")


# ── JSON ──────────────────────────────────────────────────────────────────

def export_json(data: Any, path: Path):
    path.parent.mkdir(parents=True, exist_ok=True)
    with open(path, "w", encoding="utf-8") as f:
        json.dump(data, f, indent=4, default=str)
    logger.info(f"JSON -> {path}")


# ── HTML ──────────────────────────────────────────────────────────────────

def export_html(
    baseline_metrics: List[Dict],
    ablation_metrics: List[Dict],
    dataset_report: Dict,
    model_report: Dict,
    full_metrics: Dict,
    path: Path,
):
    path.parent.mkdir(parents=True, exist_ok=True)
    now = datetime.now().strftime("%Y-%m-%d %H:%M:%S")

    def _table(rows: List[Dict], title: str) -> str:
        if not rows:
            return f"<h3>{title}</h3><p>No data available.</p>"
        flat = _flat(rows)
        headers = list(flat[0].keys())
        rows_html = "".join(
            "<tr>" + "".join(f"<td>{r.get(h,'')}</td>" for h in headers) + "</tr>"
            for r in flat
        )
        header_html = "".join(f"<th>{h}</th>" for h in headers)
        return f"""
        <h3>{title}</h3>
        <div class='table-wrap'>
        <table>
          <thead><tr>{header_html}</tr></thead>
          <tbody>{rows_html}</tbody>
        </table>
        </div>"""

    def _kv_section(data: Dict, title: str) -> str:
        rows = "".join(
            f"<tr><td><b>{k}</b></td><td>{v}</td></tr>"
            for k, v in data.items()
            if k not in {"confusion_matrix", "classification_report", "roc_curve", "pr_curve"}
        )
        return f"<h3>{title}</h3><table>{rows}</table>"

    html = f"""<!DOCTYPE html>
<html lang="en">
<head>
  <meta charset="UTF-8">
  <meta name="viewport" content="width=device-width, initial-scale=1.0">
  <title>Avian DNA Classifier — Research Report</title>
  <style>
    body {{ font-family: 'Segoe UI', sans-serif; margin: 0; padding: 0; background: #f8fafc; color: #1e293b; }}
    header {{ background: linear-gradient(135deg,#1e3a5f,#1d4ed8); color: white; padding: 2.5rem 3rem; }}
    header h1 {{ margin:0; font-size:2rem; }}
    header p  {{ margin:0.4rem 0 0; opacity:0.8; }}
    main {{ max-width: 1200px; margin: 2rem auto; padding: 0 2rem; }}
    section {{ background: white; border-radius: 12px; box-shadow: 0 2px 8px rgba(0,0,0,0.07); padding: 2rem; margin-bottom: 2rem; }}
    h2 {{ color: #1d4ed8; border-bottom: 2px solid #e2e8f0; padding-bottom: 0.5rem; }}
    h3 {{ color: #334155; margin-top: 1.5rem; }}
    .table-wrap {{ overflow-x: auto; }}
    table {{ border-collapse: collapse; width: 100%; font-size: 0.875rem; margin-top: 0.5rem; }}
    th {{ background: #1d4ed8; color: white; padding: 8px 12px; text-align: left; }}
    td {{ padding: 7px 12px; border-bottom: 1px solid #e2e8f0; }}
    tr:nth-child(even) {{ background: #f1f5f9; }}
    .badge {{ display:inline-block; padding:2px 10px; border-radius:999px; font-size:0.8rem; font-weight:600; }}
    .badge-blue {{ background:#dbeafe; color:#1d4ed8; }}
    footer {{ text-align:center; padding:2rem; color:#94a3b8; font-size:0.8rem; }}
  </style>
</head>
<body>
<header>
  <h1>🧬 Avian DNA Classification — Research Report</h1>
  <p>Hybrid DNABERT-2 + CNN + GAT + XGBoost Pipeline &nbsp;|&nbsp; Generated: {now}</p>
</header>
<main>
  <section>
    <h2>📊 Final Model Performance</h2>
    {_kv_section({k: v for k, v in full_metrics.items()
                  if k not in ("confusion_matrix","classification_report","roc_curve","pr_curve")},
                 "Full Metrics")}
  </section>

  <section>
    <h2>🔬 Baseline Model Comparison</h2>
    {_table(baseline_metrics, "Comparison Table")}
  </section>

  <section>
    <h2>🔩 Ablation Study</h2>
    {_table(ablation_metrics, "Component Ablation")}
  </section>

  <section>
    <h2>📁 Dataset Report</h2>
    {_kv_section(dataset_report, "Dataset Statistics")}
  </section>

  <section>
    <h2>🤖 Model Architecture</h2>
    {_kv_section(model_report, "Architecture & Hyperparameters")}
  </section>
</main>
<footer>Avian DNA Classifier Research Report &copy; 2024. For academic use.</footer>
</body>
</html>"""

    with open(path, "w", encoding="utf-8") as f:
        f.write(html)
    logger.info(f"HTML -> {path}")


# ── PDF (fpdf2) ───────────────────────────────────────────────────────────

def export_pdf(
    baseline_metrics: List[Dict],
    ablation_metrics: List[Dict],
    full_metrics: Dict,
    path: Path,
):
    try:
        from fpdf import FPDF

        pdf = FPDF()
        pdf.set_auto_page_break(auto=True, margin=15)
        pdf.add_page()
        pdf.set_font("Helvetica", "B", 16)
        pdf.cell(0, 10, "Avian DNA Classification - Research Report", ln=True, align="C")
        pdf.set_font("Helvetica", "", 10)
        pdf.cell(0, 8, f"Generated: {datetime.now().strftime('%Y-%m-%d %H:%M')}", ln=True, align="C")
        pdf.ln(6)

        def section(title):
            pdf.set_font("Helvetica", "B", 12)
            pdf.set_fill_color(29, 78, 216)
            pdf.set_text_color(255, 255, 255)
            pdf.cell(0, 8, title, ln=True, fill=True)
            pdf.set_text_color(0, 0, 0)
            pdf.set_font("Helvetica", "", 9)
            pdf.ln(2)

        def kv(k, v):
            pdf.set_font("Helvetica", "B", 9)
            pdf.cell(70, 6, str(k)[:40] + ":", ln=False)
            pdf.set_font("Helvetica", "", 9)
            pdf.cell(0, 6, str(v)[:80], ln=True)

        # Full metrics
        section("Final Model Performance")
        skip = {"confusion_matrix", "classification_report", "roc_curve", "pr_curve", "annotated_papers"}
        for k, v in full_metrics.items():
            if k not in skip:
                kv(k, v)
        pdf.ln(4)

        # Baseline table
        section("Baseline Model Comparison")
        flat = _flat(baseline_metrics)
        if flat:
            cols = ["model_name", "accuracy", "f1_score", "roc_auc", "mcc", "training_time_s"]
            col_w = [55, 22, 22, 22, 22, 30]
            pdf.set_font("Helvetica", "B", 8)
            for c, w in zip(cols, col_w):
                pdf.cell(w, 6, c.replace("_", " ").title(), border=1)
            pdf.ln()
            pdf.set_font("Helvetica", "", 8)
            for row in flat:
                for c, w in zip(cols, col_w):
                    pdf.cell(w, 6, str(row.get(c, ""))[:20], border=1)
                pdf.ln()
        pdf.ln(4)

        # Ablation table
        section("Ablation Study")
        flat_abl = _flat(ablation_metrics)
        if flat_abl:
            cols = ["model_name", "accuracy", "f1_score", "roc_auc"]
            col_w = [70, 30, 30, 30]
            pdf.set_font("Helvetica", "B", 8)
            for c, w in zip(cols, col_w):
                pdf.cell(w, 6, c.replace("_", " ").title(), border=1)
            pdf.ln()
            pdf.set_font("Helvetica", "", 8)
            for row in flat_abl:
                for c, w in zip(cols, col_w):
                    pdf.cell(w, 6, str(row.get(c, ""))[:25], border=1)
                pdf.ln()

        path.parent.mkdir(parents=True, exist_ok=True)
        pdf.output(str(path))
        logger.info(f"PDF -> {path}")

    except ImportError:
        logger.warning("fpdf2 not installed. Skipping PDF export. Run: pip install fpdf2")
    except Exception as e:
        logger.error(f"PDF export failed: {e}")


# ── Master Export ─────────────────────────────────────────────────────────

def export_all(
    baseline_metrics: List[Dict],
    ablation_metrics: List[Dict],
    full_metrics: Dict,
    dataset_report: Dict,
    model_report: Dict,
    output_dir: Path,
):
    """Run all export formats in one call."""
    reports_dir = output_dir / "reports"
    tables_dir  = output_dir / "tables"

    flat_base = _flat(baseline_metrics)
    flat_abl  = _flat(ablation_metrics)

    export_csv(flat_base, tables_dir / "baseline_comparison.csv")
    export_csv(flat_abl,  tables_dir / "ablation_study.csv")

    export_json({"baseline": baseline_metrics, "ablation": ablation_metrics,
                 "full_metrics": full_metrics, "dataset": dataset_report,
                 "model": model_report}, reports_dir / "results.json")

    export_excel({
        "Baseline Comparison": flat_base,
        "Ablation Study":      flat_abl,
        "Dataset Report":      [dataset_report],
        "Model Report":        [model_report],
    }, reports_dir / "results.xlsx")

    export_html(
        baseline_metrics, ablation_metrics,
        dataset_report, model_report, full_metrics,
        reports_dir / "report.html",
    )

    export_pdf(
        baseline_metrics, ablation_metrics, full_metrics,
        reports_dir / "report.pdf",
    )

    logger.info(f"All exports complete → {output_dir}")
